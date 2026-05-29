from __future__ import annotations

import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches
from sqlalchemy.orm import Session

from app.firms.matcher import FIRMS_CAVEAT
from app.models import Case, CaseMatch, EvidenceFile, Export, FirmsPoint, Message


def _export_dir() -> Path:
    path = Path("data/exports") / datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    path.mkdir(parents=True, exist_ok=True)
    return path


def _case_matches(session: Session, case_id: int) -> list[CaseMatch]:
    return session.query(CaseMatch).filter(CaseMatch.case_id == case_id).order_by(CaseMatch.total_score.desc()).all()


def _evidence_files(session: Session, case_id: int) -> list[EvidenceFile]:
    return session.query(EvidenceFile).filter(EvidenceFile.case_id == case_id).all()


def build_docx_report(session: Session, out_path: Path) -> None:
    doc = Document()
    doc.add_heading("Striker evidence report", 0)
    doc.add_paragraph("Retrospective evidence package. No live tracking, predictions, routes, targets, or tactical conclusions.")
    cases = session.query(Case).order_by(Case.id).all()
    table = doc.add_table(rows=1, cols=5)
    for idx, title in enumerate(["Case", "Date", "Place", "Coordinates", "Matches"]):
        table.rows[0].cells[idx].text = title
    for case in cases:
        row = table.add_row().cells
        row[0].text = str(case.id)
        row[1].text = case.event_date.isoformat() if case.event_date else ""
        row[2].text = case.place_name or ""
        row[3].text = f"{case.lat}, {case.lon}" if case.lat is not None and case.lon is not None else ""
        row[4].text = str(len(_case_matches(session, case.id)))
    for case in cases:
        doc.add_page_break()
        doc.add_heading(f"Case {case.id}", level=1)
        doc.add_paragraph(case.raw_text)
        doc.add_paragraph(f"Date/time/place: {case.event_date or ''} {case.event_time_local or ''} {case.place_name or ''}")
        matches = _case_matches(session, case.id)
        mt = doc.add_table(rows=1, cols=6)
        for idx, title in enumerate(["Type", "Priority", "Total", "Time", "Geo", "Explanation"]):
            mt.rows[0].cells[idx].text = title
        for match in matches:
            row = mt.add_row().cells
            row[0].text = match.match_type
            row[1].text = match.priority
            row[2].text = str(match.total_score)
            row[3].text = str(match.time_score)
            row[4].text = str(match.geo_score)
            row[5].text = match.explanation
            if match.message_id:
                message = session.get(Message, match.message_id)
                if message and message.url:
                    doc.add_paragraph(f"Telegram link: {message.url}")
        doc.add_paragraph(FIRMS_CAVEAT)
        for evidence in _evidence_files(session, case.id):
            if evidence.path.lower().endswith(".png") and Path(evidence.path).exists():
                try:
                    doc.add_picture(evidence.path, width=Inches(6.5))
                except Exception:
                    doc.add_paragraph(f"Evidence file: {evidence.path}")
    doc.save(out_path)


def build_html_report(session: Session, out_path: Path) -> None:
    parts = [
        "<!doctype html><html><head><meta charset='utf-8'><title>Striker evidence report</title>",
        "<style>body{font-family:Arial,sans-serif;margin:32px;line-height:1.45} table{border-collapse:collapse;width:100%;margin:12px 0} td,th{border:1px solid #ccd2da;padding:6px;vertical-align:top} .caveat{background:#fff8dc;padding:10px;border:1px solid #ead58b}</style>",
        "</head><body><h1>Striker evidence report</h1>",
        "<p>Retrospective evidence package. No live tracking, predictions, routes, targets, or tactical conclusions.</p>",
    ]
    for case in session.query(Case).order_by(Case.id).all():
        parts.append(f"<h2>Case {case.id}</h2><p>{case.raw_text}</p>")
        parts.append(f"<p>{case.event_date or ''} {case.event_time_local or ''} {case.place_name or ''}</p>")
        parts.append("<table><tr><th>Type</th><th>Priority</th><th>Total</th><th>Explanation</th><th>Link</th></tr>")
        for match in _case_matches(session, case.id):
            link = ""
            if match.message_id:
                message = session.get(Message, match.message_id)
                if message and message.url:
                    link = f"<a href='{message.url}'>{message.url}</a>"
            parts.append(
                f"<tr><td>{match.match_type}</td><td>{match.priority}</td><td>{match.total_score}</td><td>{match.explanation}</td><td>{link}</td></tr>"
            )
        parts.append("</table>")
        parts.append(f"<p class='caveat'>{FIRMS_CAVEAT}</p>")
    parts.append("</body></html>")
    out_path.write_text("\n".join(parts), encoding="utf-8")


def build_zip(session: Session, export_dir: Path, docx_path: Path, html_path: Path, source_docx: str | None) -> Path:
    zip_path = export_dir / "evidence.zip"
    metadata = {
        "created_at": datetime.utcnow().isoformat(),
        "cases": [case.id for case in session.query(Case).all()],
        "caveat": FIRMS_CAVEAT,
    }
    (export_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    if source_docx and Path(source_docx).exists():
        source_dir = export_dir / "source_docx_copy"
        source_dir.mkdir(exist_ok=True)
        shutil.copy2(source_docx, source_dir / Path(source_docx).name)
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for file in [docx_path, html_path, export_dir / "metadata.json"]:
            zf.write(file, file.name)
        for folder in ["data/screenshots", "data/maps"]:
            base = Path(folder)
            if base.exists():
                for path in base.glob("*"):
                    if path.is_file():
                        zf.write(path, f"{base.name}/{path.name}")
        source_dir = export_dir / "source_docx_copy"
        if source_dir.exists():
            for path in source_dir.glob("*"):
                zf.write(path, f"source_docx_copy/{path.name}")
    return zip_path


def export_report(session: Session, job_id: int | None = None, source_docx: str | None = None) -> Export:
    export_dir = _export_dir()
    docx_path = export_dir / "report.docx"
    html_path = export_dir / "report.html"
    build_docx_report(session, docx_path)
    build_html_report(session, html_path)
    zip_path = build_zip(session, export_dir, docx_path, html_path, source_docx)
    export = Export(
        job_id=job_id,
        title="Striker evidence report",
        docx_path=str(docx_path),
        html_path=str(html_path),
        zip_path=str(zip_path),
    )
    session.add(export)
    session.flush()
    return export

