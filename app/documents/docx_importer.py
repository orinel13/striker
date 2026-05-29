from __future__ import annotations

import json
from pathlib import Path
from datetime import date

from docx import Document
from sqlalchemy.orm import Session

from app.config import get_settings
from app.documents.strike_table_parser import ParsedStrikeRow, parse_docx_strike_rows
from app.geo.gazetteer import Gazetteer
from app.models import Case
from app.nlp.extractors import build_time_window, extract_coordinates, extract_date, extract_time


def iter_docx_blocks(path: Path | str) -> list[str]:
    doc = Document(str(path))
    blocks: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in doc.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                blocks.append(" | ".join(values))
    return blocks


def _case_from_parsed(session: Session, parsed: ParsedStrikeRow, source: str, gazetteer: Gazetteer) -> Case | None:
    if not any([parsed.event_date, parsed.event_time_local, parsed.time_range_start, parsed.place_name_raw, parsed.lat is not None, parsed.lon is not None]):
        return None
    lat = parsed.lat
    lon = parsed.lon
    coordinate_source = parsed.coordinate_source
    if (lat is None or lon is None) and parsed.place_name_raw:
        place_match = gazetteer.find(parsed.place_name_raw)
        if place_match:
            lat = place_match.place.lat
            lon = place_match.place.lon
            coordinate_source = "gazetteer"
    case = Case(
        source_docx=source,
        raw_text=parsed.raw_text,
        event_date=parsed.event_date,
        event_time_local=parsed.event_time_local or parsed.time_range_start,
        time_window_start=parsed.time_window_start,
        time_window_end=parsed.time_window_end,
        place_name=parsed.place_name_raw,
        oblast=parsed.oblast,
        reference_text=parsed.reference_text,
        raw_grid_northing=parsed.raw_grid_northing,
        raw_grid_easting=parsed.raw_grid_easting,
        coordinate_source=coordinate_source,
        parser_warnings=json.dumps(parsed.parser_warnings, ensure_ascii=False) if parsed.parser_warnings else None,
        lat=lat,
        lon=lon,
        radius_km=get_settings().firms_default_radius_km,
        attack_type="unknown",
    )
    session.add(case)
    return case


def import_docx(
    session: Session,
    path: Path | str,
    document_date: date | None = None,
    default_year: int | None = None,
    period_start: date | None = None,
    period_end: date | None = None,
    night_mode: bool = False,
    rollover_hour: int = 12,
) -> list[Case]:
    settings = get_settings()
    source = str(path)
    gazetteer = Gazetteer(session)
    cases: list[Case] = []
    parsed_rows = parse_docx_strike_rows(
        path,
        document_date=document_date,
        default_year=default_year,
        period_start=period_start,
        period_end=period_end,
        night_mode=night_mode,
        rollover_hour=rollover_hour,
    )
    if parsed_rows:
        for parsed in parsed_rows:
            case = _case_from_parsed(session, parsed, source, gazetteer)
            if case:
                cases.append(case)
        session.flush()
        return cases
    for block in iter_docx_blocks(path):
        event_date = extract_date(block)
        event_time = extract_time(block)
        coords = extract_coordinates(block)
        place_match = gazetteer.find(block)
        lat = coords[0] if coords else (place_match.place.lat if place_match else None)
        lon = coords[1] if coords else (place_match.place.lon if place_match else None)
        place_name = place_match.place.name if place_match else None
        window_start, window_end = build_time_window(event_date, event_time)
        if not any([event_date, coords, place_match]):
            continue
        case = Case(
            source_docx=source,
            raw_text=block,
            event_date=event_date,
            event_time_local=event_time,
            time_window_start=window_start,
            time_window_end=window_end,
            place_name=place_name,
            lat=lat,
            lon=lon,
            radius_km=settings.firms_default_radius_km,
            attack_type="unknown",
        )
        session.add(case)
        cases.append(case)
    session.flush()
    return cases
