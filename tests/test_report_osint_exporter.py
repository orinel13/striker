from datetime import datetime

from docx import Document

from app.documents import report_exporter
from app.documents.report_exporter import build_osint_docx_report
from app.models import Case, CaseMatch, Channel, Message
from app.telegram.normalizer import normalize_text, text_sha256


def _message(session, channel, tg_id: int, text: str, url: str, posted_at: datetime) -> Message:
    norm = normalize_text(text)
    message = Message(
        channel_id=channel.id,
        tg_message_id=tg_id,
        posted_at=posted_at,
        collected_at=posted_at,
        text=text,
        normalized_text=norm,
        text_hash=text_sha256(norm),
        url=url,
    )
    session.add(message)
    session.flush()
    return message


def test_osint_docx_report_human_structure(session, tmp_path, monkeypatch):
    monkeypatch.setattr(report_exporter, "ensure_publication_screenshot", lambda session, item: None)
    ch1 = Channel(username="sloviansk", title="Славянск Online")
    ch2 = Channel(username="kramatorsk", title="Краматорск")
    session.add_all([ch1, ch2])
    session.flush()
    c1 = Case(raw_text="case slov 1", place_name="Славянск")
    c2 = Case(raw_text="case slov 2", place_name="Славянск")
    c3 = Case(raw_text="case kram", place_name="Краматорск")
    c4 = Case(raw_text="case empty", place_name="Дружковка")
    session.add_all([c1, c2, c3, c4])
    session.flush()
    m1 = _message(session, ch1, 1, "публикация 1", "https://t.me/sloviansk/1", datetime(2024, 5, 29, 7, 0))
    m2 = _message(session, ch1, 2, "публикация 2", "https://t.me/sloviansk/2", datetime(2024, 5, 29, 8, 0))
    m3 = _message(session, ch2, 1, "публикация 3", "https://t.me/kramatorsk/1", datetime(2024, 5, 29, 9, 0))
    for case, message in [(c1, m1), (c2, m2), (c3, m3)]:
        session.add(
            CaseMatch(
                case_id=case.id,
                message_id=message.id,
                match_type="telegram",
                priority="A",
                total_score=0.9,
                explanation="match",
            )
        )
    session.add(CaseMatch(case_id=c4.id, match_type="none", priority="NO DATA", explanation="none"))
    session.commit()
    out = tmp_path / "report.docx"
    build_osint_docx_report(session, out)
    parsed = Document(out)
    text_parts = [p.text for p in parsed.paragraphs]
    for table in parsed.tables:
        for row in table.rows:
            text_parts.extend(cell.text for cell in row.cells)
    text = "\n".join(text_parts)
    assert "Славянск" in text
    assert "Краматорск" in text
    assert "Публикация №1" in text
    assert "Публикация №2" in text
    assert "Публикация №3" in text
    assert "Case 1" not in text
    assert "Field Value" not in text
    assert "NO DATA" not in text
    assert "https://t.me/sloviansk/1" in text
    assert "Дружковка" in text
    assert "Публикация №4" not in text
